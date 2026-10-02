# -*- coding: utf-8 -*-
"""
航班行程安排助手 v6
- 调机合并：跳过新需求前的绕路原调机，直接从最后非调机段出发地飞到新需求出发地
- 回程调机：新需求完成后，自动接回原计划目的地
- 保留：飞行时间自学习、多飞机对比、放宽模式、原/新调机标注
"""

import streamlit as st
import re
import io
from math import radians, sin, cos, asin, sqrt
from datetime import datetime, timedelta, date
from openpyxl import load_workbook

st.set_page_config(page_title="航班行程安排助手", page_icon="✈️", layout="wide")
st.title("✈️ 航班行程安排助手")
st.caption("上传航段数据 → 自动选最优 / 指定飞机 → 排行程 · 推荐起飞时刻 · 调机合并")


# ================================================================
# 城市 → 四字码
# ================================================================
CITY_TO_ICAO = {
    "澳门": "VMMC", "香港": "VHHH", "台北松山": "RCSS", "台北桃园": "RCTP",
    "北京大兴": "ZBAD", "北京首都": "ZBAA",
    "上海虹桥": "ZSSS", "上海浦东": "ZSPD",
    "成都双流": "ZUUU", "成都天府": "ZUTF",
    "深圳宝安": "ZGSZ", "广州白云": "ZGGG",
    "宁波栎社": "ZSNB", "温州龙湾": "ZSWZ",
    "厦门高崎": "ZSAM", "泉州晋江": "ZSQZ",
    "杭州萧山": "ZSHC", "南京禄口": "ZSNJ",
    "天津滨海": "ZBTJ", "沈阳桃仙": "ZYTX",
    "乌鲁木齐天山": "ZWWW", "重庆江北": "ZUCK",
    "西宁曹家堡": "ZLXN", "鄂尔多斯伊金霍洛": "ZBDS",
    "丽江三义": "ZPLJ", "威海大水泊": "ZSWH",
    "长沙黄花": "ZGHA", "嘉兴南湖": "ZSJX",
    "宜昌三峡": "ZHYC", "济宁大安": "ZSJG",
    "福州长乐": "ZSFZ", "青岛胶东": "ZSQD",
    "武汉天河": "ZHHH", "海口美兰": "ZJHK",
    "三亚凤凰": "ZJSY", "珠海金湾": "ZGSD",
    "大连周水子": "ZYTL", "合肥新桥": "ZSOF",
    "南昌昌北": "ZSCN", "贵阳龙洞堡": "ZUGY",
    "桂林两江": "ZGKL", "哈尔滨太平": "ZYHB",
    "长春龙嘉": "ZYCC", "台中清泉岗": "RCMQ",
    "越南岘港": "VVDN", "柬埔寨金边 德崇": "VDTI", "柬埔寨金边": "VDTI",
    "老挝万象": "VLVT", "瓦岱": "VLVT",
    "日本东京 羽田": "RJTT", "日本东京 成田": "RJAA",
    "日本大阪 关西": "RJBB", "日本那霸": "ROAH",
    "韩国济州岛": "RKPC", "韩国首尔": "RKSI",
    "新加坡实里达": "WSSL", "新加坡樟宜": "WSSS",
    "泰国曼谷 廊曼": "VTBD", "泰国普吉": "VTSP",
    "泰国清莱": "VTCT", "菲律宾马尼拉": "RPLL",
    "印尼万鸦老": "WAMM", "印尼韦达港": "WAEH",
    "印尼雅加达 哈达": "WIII", "印尼龙目岛普拉亚": "WADL",
    "澳大利亚悉尼 金斯福德": "YSSY",
    "澳大利亚墨尔本": "YMML",
    "澳大利亚朗塞斯顿": "YMLT",
    "加拿大蒙特利尔 特鲁多": "CYUL",
    "加拿大温哥华": "CYVR",
    "美国贝德福德": "KBED", "美国达拉斯": "KDFW",
    "美国费城": "KPHL", "美国旧金山": "KSFO",
    "美国加利福尼亚州圣迭戈": "KSDM",
    "马尔代夫马法鲁岛": "VRDA",
    "捷克布拉格 鲁济涅": "LKPR",
    "法国巴黎 布尔歇": "LFPB",
    "西班牙马德里 巴拉哈斯": "LEMD",
    "罗马": "LIRA", "埃及卢克索": "HELX",
    "埃及阿布辛伯勒": "HEBL", "埃及古尔代盖": "HEGN",
    "埃及开罗": "HECA", "文莱斯里巴加湾": "WBSB",
    "厄瓜多尔基多": "SEQM", "秘鲁库斯科": "SPZO",
    "玻利维亚乌尤尼": "SLUY", "智利伊基克": "SCDA",
    "智利卡拉马": "SCCF", "秘鲁利马": "SPJC",
    "巴哈马拿骚": "MYNN",
}


def get_icao(city):
    if not city:
        return None
    city = str(city).strip()
    if city in CITY_TO_ICAO:
        return CITY_TO_ICAO[city]
    for k, v in CITY_TO_ICAO.items():
        if k in city or city in k:
            return v
    return None


def icao_to_city(icao):
    for k, v in CITY_TO_ICAO.items():
        if v == icao:
            return k
    return icao


# ================================================================
# 机场坐标
# ================================================================
AIRPORT_COORDS = {
    "ZBAD": (39.509, 116.411), "ZBAA": (40.079, 116.585),
    "ZSSS": (31.197, 121.336), "ZSPD": (31.143, 121.805),
    "ZUUU": (30.578, 103.947), "ZUTF": (30.312, 104.441),
    "ZGSZ": (22.639, 113.811), "ZGGG": (23.392, 113.298),
    "ZSNB": (29.826, 121.462), "ZSWZ": (27.912, 120.852),
    "ZSAM": (24.544, 118.128), "ZSQZ": (24.796, 118.590),
    "ZSHC": (30.229, 120.434), "ZSNJ": (31.742, 118.862),
    "ZBTJ": (39.124, 117.346), "ZYTX": (41.639, 123.483),
    "ZWWW": (43.907, 87.474), "ZUCK": (29.719, 106.641),
    "ZLXN": (36.528, 102.043), "ZBDS": (39.494, 109.859),
    "ZPLJ": (26.680, 100.246), "ZSWH": (37.187, 122.229),
    "ZGHA": (28.189, 113.219), "ZSJX": (30.701, 120.665),
    "ZHYC": (30.552, 111.479), "ZSJG": (35.496, 116.740),
    "ZSFZ": (25.935, 119.663), "ZSQD": (36.361, 120.086),
    "ZHHH": (30.784, 114.208), "ZJHK": (19.935, 110.459),
    "ZJSY": (18.303, 109.412), "ZGSD": (22.006, 113.376),
    "ZYTL": (38.966, 121.538), "ZSOF": (31.780, 116.976),
    "ZSCN": (28.865, 115.900), "ZUGY": (26.539, 106.801),
    "ZGKL": (25.218, 110.039), "ZYHB": (45.623, 126.250),
    "ZYCC": (43.996, 125.685),
    "VMMC": (22.149, 113.592), "VHHH": (22.309, 113.915),
    "RCSS": (25.069, 121.552), "RCTP": (25.077, 121.233),
    "RCMQ": (24.264, 120.621),
    "VVDN": (16.044, 108.199), "VDTI": (11.547, 104.844),
    "VLVT": (17.988, 102.563),
    "RJTT": (35.549, 139.779), "RJAA": (35.765, 140.386),
    "RJBB": (34.434, 135.232), "ROAH": (26.196, 127.646),
    "RKPC": (33.512, 126.493), "RKSI": (37.469, 126.451),
    "WSSL": (1.417, 103.868), "WSSS": (1.357, 103.988),
    "VTBD": (13.912, 100.607), "VTSP": (8.113, 98.317),
    "VTCT": (19.953, 99.883), "RPLL": (14.508, 121.020),
    "WAMM": (1.549, 124.926), "WAEH": (-0.836, 127.786),
    "WIII": (-6.125, 106.656), "WADL": (-8.757, 116.276),
    "YSSY": (-33.946, 151.177), "YMML": (-37.673, 144.843),
    "YMLT": (-41.545, 147.214),
    "CYUL": (45.470, -73.741), "CYVR": (49.196, -123.182),
    "KBED": (42.470, -71.289), "KDFW": (32.897, -97.038),
    "KPHL": (39.872, -75.241), "KSFO": (37.619, -122.375),
    "KSDM": (32.572, -116.980),
    "VRDA": (4.191, 73.500),
    "LKPR": (50.101, 14.260), "LFPB": (48.969, 2.441),
    "LEMD": (40.472, -3.561), "LIRA": (41.799, 12.594),
    "HELX": (25.671, 32.707), "HEBL": (22.376, 31.611),
    "HEGN": (27.178, 33.799), "HECA": (30.111, 31.413),
    "WBSB": (4.944, 114.928),
    "SEQM": (-0.129, -78.358), "SPZO": (-13.536, -71.939),
    "SLUY": (-20.442, -66.848), "SCDA": (-20.535, -70.181),
    "SCCF": (-22.498, -68.904), "SPJC": (-12.022, -77.114),
    "MYNN": (25.039, -77.466),
}


def haversine_km(lat1, lon1, lat2, lon2):
    R = 6371
    dlat = radians(lat2 - lat1)
    dlon = radians(lon2 - lon1)
    a = sin(dlat / 2) ** 2 + cos(radians(lat1)) * cos(radians(lat2)) * sin(dlon / 2) ** 2
    return 2 * R * asin(sqrt(a))


def _airport_coord_fuzzy(icao):
    if icao in AIRPORT_COORDS:
        return AIRPORT_COORDS[icao]
    prefix = icao[:1] if icao else ""
    same_region = [c for c in AIRPORT_COORDS if c.startswith(prefix)]
    if same_region:
        return AIRPORT_COORDS[same_region[0]]
    return None


# ================================================================
# 飞行时间
# ================================================================
LEARNED_FLIGHT_MIN = {}


def estimate_distance_km(from_icao, to_icao):
    if not from_icao or not to_icao:
        return None
    cf = _airport_coord_fuzzy(from_icao)
    ct = _airport_coord_fuzzy(to_icao)
    if not cf or not ct:
        return None
    return round(haversine_km(cf[0], cf[1], ct[0], ct[1]))


def estimate_flight_minutes(from_icao, to_icao):
    if not from_icao or not to_icao:
        return None
    if from_icao == to_icao:
        return 0
    learned = LEARNED_FLIGHT_MIN.get((from_icao, to_icao))
    if learned:
        return round(sum(learned) / len(learned))
    cf = _airport_coord_fuzzy(from_icao)
    ct = _airport_coord_fuzzy(to_icao)
    if not cf or not ct:
        return None
    d = haversine_km(cf[0], cf[1], ct[0], ct[1])
    if d < 500:
        speed, overhead = 650, 20
    elif d < 1500:
        speed, overhead = 720, 28
    elif d < 3000:
        speed, overhead = 760, 35
    else:
        speed, overhead = 780, 45
    return round(d / speed * 60 + overhead)


def flight_min_source(from_icao, to_icao):
    if (from_icao, to_icao) in LEARNED_FLIGHT_MIN:
        return "Excel 实测均值"
    if (from_icao in AIRPORT_COORDS) and (to_icao in AIRPORT_COORDS):
        return "公式估算"
    return "公式估算（坐标模糊）"


# ================================================================
# 时间工具
# ================================================================
def time_to_min(t):
    if t is None:
        return None
    if isinstance(t, datetime):
        return t.hour * 60 + t.minute
    try:
        if hasattr(t, "hour") and hasattr(t, "minute"):
            return t.hour * 60 + t.minute
    except Exception:
        pass
    m = re.match(r"^(\d{1,2}):(\d{2})$", str(t).strip())
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    return None


def min_to_time_str(m):
    m = m % 1440
    return f"{m // 60:02d}:{m % 60:02d}"


def min_to_dur_str(m):
    m = int(m)
    sign = "-" if m < 0 else ""
    m = abs(m)
    h, mi = divmod(m, 60)
    if h and mi:
        return f"{sign}{h}h{mi:02d}m"
    elif h:
        return f"{sign}{h}h"
    return f"{sign}{mi}m"


def plan_minutes(flight_minutes):
    total = flight_minutes + 15
    return ((total + 4) // 5) * 5


def is_domestic(icao):
    return bool(icao) and str(icao).startswith("Z")


def transit_minutes(icao):
    return 90 if is_domestic(icao) else 120


def is_ferry_use(use_text):
    if not use_text:
        return False
    t = str(use_text)
    return ("调机" in t) or ("维修" in t)


def to_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


# ================================================================
# 解析 Excel
# ================================================================
def load_excel_plan(file_bytes):
    global LEARNED_FLIGHT_MIN
    LEARNED_FLIGHT_MIN = {}
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb["航段(北京时)"] if "航段(北京时)" in wb.sheetnames else wb.active

    rows = []
    for r in range(3, ws.max_row + 1):
        reg = ws.cell(r, 3).value
        if not reg:
            continue
        reg_str = str(reg).strip().upper()
        dep_date = to_date(ws.cell(r, 7).value)
        arr_date = to_date(ws.cell(r, 15).value)
        dep_icao = str(ws.cell(r, 11).value or "").strip().upper()
        dep_city = str(ws.cell(r, 12).value or "").strip()
        arr_icao = str(ws.cell(r, 13).value or "").strip().upper()
        arr_city = str(ws.cell(r, 14).value or "").strip()
        dep_city_clean = dep_city.split()[0] if dep_city else ""
        arr_city_clean = arr_city.split()[0] if arr_city else ""
        dep_min = time_to_min(ws.cell(r, 8).value)
        arr_min = time_to_min(ws.cell(r, 16).value)
        if dep_min is None or arr_min is None or not dep_date:
            continue
        arr_dt_date = arr_date or dep_date
        if arr_date is None and arr_min < dep_min:
            arr_dt_date = dep_date + timedelta(days=1)

        sched_min = time_to_min(ws.cell(r, 19).value)
        if sched_min and dep_icao and arr_icao and dep_icao != arr_icao:
            LEARNED_FLIGHT_MIN.setdefault((dep_icao, arr_icao), []).append(sched_min)
            LEARNED_FLIGHT_MIN.setdefault((arr_icao, dep_icao), []).append(sched_min)

        rows.append({
            "reg": reg_str,
            "use": str(ws.cell(r, 4).value or "").strip(),
            "dep_date": dep_date,
            "dep_min": dep_min,
            "dep_icao": dep_icao,
            "dep_city": dep_city_clean,
            "arr_date": arr_dt_date,
            "arr_min": arr_min,
            "arr_icao": arr_icao,
            "arr_city": arr_city_clean,
        })
    return rows


# ================================================================
# 解析需求
# ================================================================
REL_DATE_KW = {
    "今天": 0, "今日": 0, "明天": 1, "明日": 1,
    "后天": 2, "後天": 2, "大后天": 3,
}
TIME_WORDS = ["上午", "下午", "晚上", "早上", "凌晨", "中午", "傍晚", "夜里", "晚间"]


def parse_request(text):
    clean_text = text
    today = datetime.now().date()
    target_date = None
    for kw, delta in REL_DATE_KW.items():
        if kw in clean_text:
            target_date = today + timedelta(days=delta)
            break
    if target_date is None:
        m = re.search(r"(\d{1,2})\s*[.\-/月]\s*(\d{1,2})\s*[日号]?", clean_text)
        if m:
            mon, day = int(m.group(1)), int(m.group(2))
            year = today.year
            try:
                d = date(year, mon, day)
                if (today - d).days > 180:
                    d = date(year + 1, mon, day)
                target_date = d
            except Exception:
                pass

    route_text = clean_text
    for kw in list(REL_DATE_KW.keys()) + TIME_WORDS:
        route_text = route_text.replace(kw, " ")
    route_text = re.sub(r"\d{1,2}\s*[.\-/月]\s*\d{1,2}\s*[日号]?", " ", route_text)

    route = None
    m = re.search(
        r"([\u4e00-\u9fff][\u4e00-\u9fff\s]{0,20}?)\s*[-–—到至]\s*"
        r"([\u4e00-\u9fff][\u4e00-\u9fff\s]{0,20})",
        route_text,
    )
    if m:
        dep_c = m.group(1).strip()
        arr_c = m.group(2).strip()
        if dep_c and arr_c:
            route = (dep_c, arr_c)

    flight_min = None
    has_explicit_time = False
    for c in re.findall(r"(?<!\d)(\d{4})(?!\d)", clean_text):
        hh, mm = int(c[:2]), int(c[2:])
        if hh < 24 and mm < 60:
            flight_min = hh * 60 + mm
            has_explicit_time = True
            break
    if flight_min is None:
        m = re.search(
            r"(\d{1,2})\s*(?:h|H|小时)\s*(?:(\d{1,2})\s*(?:m|M|分钟?))?",
            clean_text,
        )
        if m:
            hh = int(m.group(1))
            mm = int(m.group(2)) if m.group(2) else 0
            if hh < 24 and mm < 60:
                flight_min = hh * 60 + mm
                has_explicit_time = True
    if flight_min is None:
        m = re.search(r"(?<!\d)(\d{1,2}):(\d{2})(?!\d)", clean_text)
        if m:
            hh, mm = int(m.group(1)), int(m.group(2))
            if hh < 24 and mm < 60:
                flight_min = hh * 60 + mm
                has_explicit_time = True

    return {
        "date": target_date, "route": route,
        "flight_min": flight_min, "has_explicit_time": has_explicit_time,
    }


# ================================================================
# 分析（★ 新增 merge_ferry 参数）
# ================================================================
DUTY_MAX_MIN = 14 * 60
FLIGHT_MAX_MIN = 10 * 60
NEW_FERRY_TAGS = {"新增调机", "调机(优化)", "调机(合并)", "调机(回程)"}
ORIG_FERRY_TAGS = {"原调机"}


def analyze_aircraft(flights_all, reg, target_date, dep_city, arr_city,
                     flight_min, transit, plan_min, forced_main_dep=None,
                     merge_ferry=False):
    dep_icao_req = get_icao(dep_city)
    arr_icao_req = get_icao(arr_city)
    if not dep_icao_req:
        return {"error": f"❌ 无法识别出发城市：{dep_city}"}
    if not arr_icao_req:
        return {"error": f"❌ 无法识别到达城市：{arr_city}"}

    reg_flights = [f for f in flights_all if f["reg"] == reg]
    reg_flights.sort(key=lambda x: (x["dep_date"], x["dep_min"]))
    before = [f for f in reg_flights if f["dep_date"] <= target_date]
    day_flights = [f for f in reg_flights if f["dep_date"] == target_date]
    if not before:
        return {"error": f"❌ 飞机 {reg} 在 {target_date.strftime('%m月%d日')} 之前没有任何航段。"}

    # ---------- 调机合并：找"最后非调机段"和"后续原调机段" ----------
    last_pax = None
    later_ferry = []
    for f in day_flights:
        if is_ferry_use(f.get("use", "")):
            if last_pax is not None:
                later_ferry.append(f)
        else:
            last_pax = f
            later_ferry = []

    merged_info = None
    ferry_info = None

    if merge_ferry and last_pax and later_ferry:
        # 原调机段的最终目的地
        final_ferry_dest_icao = later_ferry[-1]["arr_icao"]
        final_ferry_dest_city = later_ferry[-1]["arr_city"]

        if final_ferry_dest_icao != dep_icao_req:
            # 方案 B：从 last_pax 到达地直接飞到新需求出发地
            ferry_from_icao = last_pax["arr_icao"]
            ferry_from_city = last_pax["arr_city"]

            if ferry_from_icao == dep_icao_req:
                ferry_min = 0
                ferry_plan_min = 0
                ferry_dep_min = (
                    (last_pax["arr_date"] - target_date).days * 1440
                    + last_pax["arr_min"]
                )
                ferry_arr_min = ferry_dep_min
            else:
                ferry_min = estimate_flight_minutes(ferry_from_icao, dep_icao_req)
                if ferry_min is not None:
                    ferry_plan_min = plan_minutes(ferry_min)
                    ferry_dep_min = (
                        (last_pax["arr_date"] - target_date).days * 1440
                        + last_pax["arr_min"]
                        + transit_minutes(ferry_from_icao)
                    )
                    ferry_arr_min = ferry_dep_min + ferry_plan_min

            if ferry_min is not None:
                ferry_info = {
                    "from_city": ferry_from_city,
                    "from_icao": ferry_from_icao,
                    "ferry_min": ferry_min,
                    "ferry_plan_min": ferry_plan_min,
                    "dep_min": ferry_dep_min,
                    "arr_min": ferry_arr_min,
                    "based_on": last_pax,
                }
                merged_info = {
                    "skipped": later_ferry,
                    "last_pax": last_pax,
                    "return_to_icao": final_ferry_dest_icao,
                    "return_to_city": final_ferry_dest_city,
                }

    # ---------- 方案 A（原逻辑）----------
    if ferry_info is None:
        if day_flights:
            last_of_day = day_flights[-1]
            if last_of_day["arr_icao"] == dep_icao_req:
                ferry_info = None
            else:
                ferry_from_icao = last_of_day["arr_icao"]
                ferry_from_city = last_of_day["arr_city"]
                ferry_min = estimate_flight_minutes(ferry_from_icao, dep_icao_req)
                if ferry_min is None:
                    return {"error": f"❌ 缺少坐标，无法估算 {ferry_from_icao} → {dep_icao_req} 调机时间。"}
                last_arr_abs = (
                    (last_of_day["arr_date"] - target_date).days * 1440
                    + last_of_day["arr_min"]
                )
                ferry_dep_min = last_arr_abs + transit_minutes(ferry_from_icao)
                ferry_plan_min = plan_minutes(ferry_min)
                ferry_arr_min = ferry_dep_min + ferry_plan_min
                ferry_info = {
                    "from_city": ferry_from_city,
                    "from_icao": ferry_from_icao,
                    "ferry_min": ferry_min,
                    "ferry_plan_min": ferry_plan_min,
                    "dep_min": ferry_dep_min,
                    "arr_min": ferry_arr_min,
                    "based_on": last_of_day,
                }
        else:
            last_seg = before[-1]
            last_city_icao = last_seg["arr_icao"]
            if last_city_icao != dep_icao_req:
                ferry_from_city = last_seg["arr_city"]
                ferry_from_icao = last_city_icao
                ferry_min = estimate_flight_minutes(ferry_from_icao, dep_icao_req)
                if ferry_min is None:
                    return {"error": f"❌ 缺少坐标，无法估算 {ferry_from_icao} → {dep_icao_req} 调机时间。"}
                ferry_dep_min = 8 * 60
                ferry_plan_min = plan_minutes(ferry_min)
                ferry_arr_min = ferry_dep_min + ferry_plan_min
                ferry_info = {
                    "from_city": ferry_from_city,
                    "from_icao": ferry_from_icao,
                    "ferry_min": ferry_min,
                    "ferry_plan_min": ferry_plan_min,
                    "dep_min": ferry_dep_min,
                    "arr_min": ferry_arr_min,
                    "based_on": last_seg,
                }
            else:
                ferry_dep_min = 8 * 60
                ferry_arr_min = ferry_dep_min
                ferry_info = {
                    "from_city": dep_city,
                    "from_icao": dep_icao_req,
                    "ferry_min": 0,
                    "ferry_plan_min": 0,
                    "dep_min": ferry_dep_min,
                    "arr_min": ferry_arr_min,
                    "based_on": last_seg,
                }

    # ---------- 主段最早可起飞 ----------
    if ferry_info is not None:
        earliest_main_dep = ferry_info["arr_min"] + transit
    else:
        last_arr_abs = (
            (day_flights[-1]["arr_date"] - target_date).days * 1440
            + day_flights[-1]["arr_min"]
        )
        earliest_main_dep = last_arr_abs + transit

    if forced_main_dep is not None:
        if forced_main_dep < earliest_main_dep:
            return {"error": (
                f"起飞 {min_to_time_str(forced_main_dep)} "
                f"早于最早可起飞 {min_to_time_str(earliest_main_dep)}"
            )}
        main_dep_min = forced_main_dep
    else:
        main_dep_min = earliest_main_dep
    main_arr_min = main_dep_min + plan_min

    # ---------- 值勤 ----------
    if day_flights:
        first_dep_abs = (
            (day_flights[0]["dep_date"] - target_date).days * 1440
            + day_flights[0]["dep_min"]
        )
        duty_start = first_dep_abs - 120
    elif ferry_info is not None:
        duty_start = ferry_info["dep_min"] - 120
    else:
        duty_start = main_dep_min - 120
    duty_end = main_arr_min + 60
    duty_total = duty_end - duty_start

    old_flight = 0
    for f in day_flights:
        d = f["dep_min"]
        a = f["arr_min"]
        if a < d:
            a += 1440
        old_flight += (a - d)

    ferry_flight = ferry_info["ferry_min"] if ferry_info else 0

    # ---------- 前日休息（绝对分钟）----------
    prev_day = target_date - timedelta(days=1)
    rest_ok = True
    rest_note = ""
    prev_day_flights = [f for f in reg_flights if f["dep_date"] == prev_day]
    if prev_day_flights:
        prev_last = prev_day_flights[-1]
        prev_duty_end_abs = (
            (prev_last["arr_date"] - target_date).days * 1440
            + prev_last["arr_min"] + 60
        )
        rest_hours = (duty_start - prev_duty_end_abs) / 60
        if rest_hours < 10:
            rest_ok = False
            rest_note = (
                f"前一日值勤结束 {prev_last['arr_date'].strftime('%m月%d日')} "
                f"{min_to_time_str(prev_duty_end_abs)}，"
                f"当日值勤开始 {min_to_time_str(duty_start)}，"
                f"休息仅 {rest_hours:.1f}h < 10h"
            )

    # ---------- 组装航段 ----------
    segments = []
    skipped_ids = set()
    if merged_info:
        for f in merged_info["skipped"]:
            skipped_ids.add((f["dep_date"], f["dep_min"], f["dep_icao"], f["arr_icao"]))

    for f in reg_flights:
        # 跳过被合并的原调机段
        key = (f["dep_date"], f["dep_min"], f["dep_icao"], f["arr_icao"])
        if key in skipped_ids:
            continue
        tag = ""
        if is_ferry_use(f.get("use", "")):
            tag = "原调机"
        segments.append({
            "date": f["dep_date"], "dep_city": f["dep_city"], "dep_icao": f["dep_icao"],
            "dep_min": f["dep_min"], "arr_min": f["arr_min"],
            "arr_city": f["arr_city"], "arr_icao": f["arr_icao"],
            "tag": tag, "use": f.get("use", ""),
        })

    if ferry_info and ferry_info["ferry_min"] > 0:
        fd_abs = ferry_info["dep_min"]
        fa_abs = ferry_info["arr_min"]
        tag = "调机(合并)" if merged_info else "新增调机"
        segments.append({
            "date": target_date + timedelta(days=fd_abs // 1440),
            "dep_city": ferry_info["from_city"], "dep_icao": ferry_info["from_icao"],
            "dep_min": fd_abs % 1440, "arr_min": fa_abs % 1440,
            "arr_city": dep_city, "arr_icao": dep_icao_req,
            "tag": tag, "use": "调机",
        })

    segments.append({
        "date": target_date + timedelta(days=main_dep_min // 1440),
        "dep_city": dep_city, "dep_icao": dep_icao_req,
        "dep_min": main_dep_min % 1440, "arr_min": main_arr_min % 1440,
        "arr_city": arr_city, "arr_icao": arr_icao_req,
        "tag": "新增", "use": "",
    })

    # ---------- 回程调机（合并模式）----------
    return_ferry_info = None
    if merged_info:
        return_to_icao = merged_info["return_to_icao"]
        return_to_city = merged_info["return_to_city"]
        if return_to_icao != arr_icao_req:
            ret_min = estimate_flight_minutes(arr_icao_req, return_to_icao)
            if ret_min is not None:
                ret_plan_min = plan_minutes(ret_min)
                ret_dep_abs = main_arr_min + transit_minutes(arr_icao_req)
                ret_arr_abs = ret_dep_abs + ret_plan_min
                segments.append({
                    "date": target_date + timedelta(days=ret_dep_abs // 1440),
                    "dep_city": arr_city, "dep_icao": arr_icao_req,
                    "dep_min": ret_dep_abs % 1440, "arr_min": ret_arr_abs % 1440,
                    "arr_city": return_to_city, "arr_icao": return_to_icao,
                    "tag": "调机(回程)", "use": "调机",
                })
                return_ferry_info = {
                    "from_city": arr_city, "from_icao": arr_icao_req,
                    "to_city": return_to_city, "to_icao": return_to_icao,
                    "ferry_min": ret_min, "ferry_plan_min": ret_plan_min,
                    "dep_min": ret_dep_abs, "arr_min": ret_arr_abs,
                }
                ferry_flight += ret_min

    total_flight = old_flight + ferry_flight + flight_min

    segments.sort(key=lambda x: (x["date"], x["dep_min"]))

    # ---------- 后续调机优化 ----------
    optimizations = []
    for i, s in enumerate(segments):
        seg_abs_dep = (s["date"] - target_date).days * 1440 + s["dep_min"]
        if seg_abs_dep <= main_arr_min:
            continue
        if s["tag"] in ("新增", "新增调机", "调机(优化)", "调机(合并)", "调机(回程)"):
            continue
        is_ferry_seg = (s["tag"] in ORIG_FERRY_TAGS) or is_ferry_use(s.get("use", ""))
        if not is_ferry_seg:
            continue
        orig_dep_icao = s["dep_icao"]
        if orig_dep_icao == arr_icao_req:
            break
        d_old = estimate_distance_km(orig_dep_icao, s["arr_icao"])
        d_new = estimate_distance_km(arr_icao_req, s["arr_icao"])
        if d_old is None or d_new is None:
            continue
        if d_new >= d_old * 0.85:
            continue

        new_flight_min = estimate_flight_minutes(arr_icao_req, s["arr_icao"])
        new_plan_min = plan_minutes(new_flight_min)
        old_dep_min = s["dep_min"]
        old_arr_min = s["arr_min"]
        new_arr_min = old_dep_min + new_plan_min

        optimizations.append({
            "date": s["date"],
            "old_dep_city": s["dep_city"], "old_dep_icao": orig_dep_icao,
            "old_tag": s["tag"] or "调机",
            "new_dep_city": icao_to_city(arr_icao_req), "new_dep_icao": arr_icao_req,
            "arr_city": s["arr_city"], "arr_icao": s["arr_icao"],
            "old_dep_time": min_to_time_str(old_dep_min),
            "old_arr_time": min_to_time_str(old_arr_min),
            "new_dep_time": min_to_time_str(old_dep_min),
            "new_arr_time": min_to_time_str(new_arr_min),
            "d_old": d_old, "d_new": d_new, "saving_km": d_old - d_new,
            "old_plan_min": plan_minutes(round(d_old / 800 * 60 + 25)),
            "new_plan_min": new_plan_min,
        })
        segments[i] = {
            "date": s["date"], "dep_city": icao_to_city(arr_icao_req),
            "dep_icao": arr_icao_req, "dep_min": old_dep_min,
            "arr_min": new_arr_min, "arr_city": s["arr_city"],
            "arr_icao": s["arr_icao"], "tag": "调机(优化)", "use": "调机",
        }
        break

    min_date = target_date - timedelta(days=1)
    max_date = target_date + timedelta(days=1)
    visible_segments = [s for s in segments if min_date <= s["date"] <= max_date]

    return {
        "reg": reg, "segments": visible_segments, "day_flights": day_flights,
        "duty_start": duty_start, "duty_end": duty_end, "duty_total": duty_total,
        "old_flight": old_flight, "ferry_info": ferry_info,
        "return_ferry_info": return_ferry_info,
        "merged_info": merged_info,
        "transit": transit, "main_dep": main_dep_min, "main_arr": main_arr_min,
        "new_duty_end": duty_end, "new_duty_total": duty_total,
        "new_flight_total": total_flight,
        "ok_duty": duty_total <= DUTY_MAX_MIN,
        "ok_flight": total_flight <= FLIGHT_MAX_MIN,
        "rest_ok": rest_ok, "rest_note": rest_note,
        "plan_min": plan_min, "flight_min": flight_min,
        "optimizations": optimizations,
    }


# ================================================================
# 起飞时刻求解
# ================================================================
def suggest_departure_times(flights_all, reg, target_date, dep_city, arr_city,
                            flight_min, objective="earliest", step=15,
                            window_h=12, relax=True, merge_ferry=False):
    dep_icao = get_icao(dep_city)
    arr_icao = get_icao(arr_city)
    if not dep_icao or not arr_icao:
        return {"error": "❌ 无法识别出发/到达城市"}
    transit = transit_minutes(dep_icao)
    plan_min = plan_minutes(flight_min)

    baseline = analyze_aircraft(
        flights_all, reg, target_date, dep_city, arr_city,
        flight_min, transit, plan_min, merge_ferry=merge_ferry,
    )
    if baseline.get("error"):
        return baseline

    earliest = baseline["main_dep"]
    feasible, fallback = [], []
    t = earliest
    end = earliest + window_h * 60
    while t <= end:
        r = analyze_aircraft(
            flights_all, reg, target_date, dep_city, arr_city,
            flight_min, transit, plan_min, forced_main_dep=t,
            merge_ferry=merge_ferry,
        )
        if not r.get("error"):
            if r["ok_duty"] and r["ok_flight"] and r["rest_ok"]:
                feasible.append(r)
            else:
                fallback.append(r)
        t += step

    if feasible:
        if objective == "earliest":
            picks = feasible[:3]
        elif objective == "latest":
            picks = feasible[-3:]
        else:
            feasible_sorted = sorted(feasible, key=lambda x: x["new_duty_total"])
            picks = feasible_sorted[:3]
        return {"baseline": baseline, "earliest": earliest,
                "feasible": feasible, "picks": picks,
                "objective": objective, "relaxed": False}

    if relax and fallback:
        def severity(r):
            return (max(0, r["new_duty_total"] - DUTY_MAX_MIN) * 2
                    + max(0, r["new_flight_total"] - FLIGHT_MAX_MIN))
        fallback.sort(key=severity)
        seen, picks = set(), []
        for r in fallback:
            if r["main_dep"] in seen:
                continue
            seen.add(r["main_dep"])
            picks.append(r)
            if len(picks) >= 3:
                break
        return {"baseline": baseline, "earliest": earliest,
                "feasible": [], "picks": picks,
                "objective": objective, "relaxed": True}

    return {"error": "❌ 在搜索窗口内未找到任何方案"}


# ================================================================
# 调机统计
# ================================================================
def _compute_ferry_metrics(result):
    fc, ft = 0, 0
    for s in result["segments"]:
        if s.get("tag") not in NEW_FERRY_TAGS:
            continue
        fc += 1
        d, a = s["dep_min"], s["arr_min"]
        if a < d:
            a += 1440
        ft += (a - d)
    return fc, ft


def _count_orig_ferry(result):
    cnt, total = 0, 0
    for s in result["segments"]:
        if s.get("tag") not in ORIG_FERRY_TAGS:
            continue
        cnt += 1
        d, a = s["dep_min"], s["arr_min"]
        if a < d:
            a += 1440
        total += (a - d)
    return cnt, total


def rank_aircraft_for_request(flights_all, all_regs, target_date,
                              dep_city, arr_city, flight_min,
                              relax=True, merge_ferry=False):
    dep_icao = get_icao(dep_city)
    arr_icao = get_icao(arr_city)
    if not dep_icao or not arr_icao:
        return []
    transit = transit_minutes(dep_icao)
    plan_min = plan_minutes(flight_min)

    candidates = []
    for reg in all_regs:
        res = analyze_aircraft(
            flights_all, reg, target_date, dep_city, arr_city,
            flight_min, transit, plan_min, merge_ferry=merge_ferry,
        )
        if res.get("error"):
            candidates.append({"reg": reg, "error": res["error"]})
            continue
        fc, ft = _compute_ferry_metrics(res)
        ofc, oft = _count_orig_ferry(res)
        ok_all = res["ok_duty"] and res["ok_flight"] and res["rest_ok"]
        over = (max(0, res["new_duty_total"] - DUTY_MAX_MIN) * 2
                + max(0, res["new_flight_total"] - FLIGHT_MAX_MIN))
        candidates.append({
            "reg": reg, "result": res,
            "ferry_count": fc, "ferry_min": ft,
            "orig_ferry_count": ofc, "orig_ferry_min": oft,
            "duty_total": res["new_duty_total"],
            "ok_all": ok_all, "over": over,
        })

    def sort_key(c):
        if "error" in c:
            return (1, 999, 999999, 999999)
        return (0 if c["ok_all"] else 1,
                c["ferry_count"], c["ferry_min"], c["over"])
    candidates.sort(key=sort_key)
    return candidates


# ================================================================
# UI
# ================================================================
uploaded = st.file_uploader(
    "① 上传航段数据导出 (.xlsx)", type=["xlsx"], key="xlsx_plan"
)
flights, all_regs = [], []
if uploaded is not None:
    try:
        flights = load_excel_plan(uploaded.getvalue())
        all_regs = sorted(set(f["reg"] for f in flights))
    except Exception as e:
        st.error(f"❌ 无法解析 Excel：{e}")

col_a, col_b = st.columns([1, 2])
with col_a:
    st.subheader("② 选择飞机")
    if all_regs:
        auto_select = st.checkbox(
            "🔍 自动选最优（对比所有飞机，优先调机段最少）",
            value=True, key="auto_select_best",
        )
        if auto_select:
            selected_reg = None
            st.caption(f"将对比 {len(all_regs)} 架：{', '.join(all_regs)}")
        else:
            selected_reg = st.selectbox(
                "选择飞机", options=all_regs,
                label_visibility="collapsed", key="reg_select",
            )
    else:
        st.info("请先上传 Excel")
        selected_reg, auto_select = None, False

with col_b:
    st.subheader("③ 潜在行程请求")
    request_text = st.text_area(
        "request", height=130, label_visibility="collapsed",
        placeholder=("例如：\n10.4 澳门-老挝万象 0158\n"
                     "明天上海虹桥-北京大兴"),
        key="request_input",
    )

st.write("")
c1, c2 = st.columns(2)
with c1:
    merge_mode = st.checkbox(
        "🔗 调机合并（跳过绕路的原调机，从最后非调机段出发地直飞新需求出发地，飞完再接回原计划）",
        value=True, key="merge_ferry_mode",
        help="例：宁波→澳门(原调机)→上海虹桥(新增调机) 会被合并为 宁波→上海虹桥(调机)，飞完新需求后自动加一段回程调机接回澳门",
    )
with c2:
    relax_mode = st.checkbox(
        "🔓 放宽模式（超时也出方案）",
        value=True, key="relax_mode",
    )

run = st.button("🚀 分析并推荐", type="primary", use_container_width=True)


# ================================================================
# 结果渲染
# ================================================================
def render_result(res, reg, dep_city, arr_city, compact=False):
    ok_all = res["ok_duty"] and res["ok_flight"] and res["rest_ok"]
    head_icon = "✅" if ok_all else "⚠️"
    if not compact:
        st.subheader(f"{head_icon} {reg} 完整行程")

    def _compact_line(seg):
        dep_hm = min_to_time_str(seg["dep_min"]).replace(":", "")
        arr_hm = min_to_time_str(seg["arr_min"]).replace(":", "")
        tag = f"  {seg['tag']}" if seg.get("tag") else ""
        return (f"{seg['date'].day}号 {seg['dep_city']}{dep_hm} "
                f"{arr_hm}{seg['arr_city']}{tag}")

    compact_lines = [reg]
    for seg in res["segments"]:
        compact_lines.append(_compact_line(seg))
    st.code("\n".join(compact_lines), language=None)

    # 调机统计
    fc, ft = _compute_ferry_metrics(res)
    ofc, oft = _count_orig_ferry(res)
    m1, m2 = st.columns(2)
    with m1:
        if fc:
            st.markdown(f"**新增/合并调机：** {fc} 段 ／ {min_to_dur_str(ft)}")
        else:
            st.markdown("**新增/合并调机：** 无 ✅")
    with m2:
        if ofc:
            st.markdown(f"**原计划调机：** {ofc} 段 ／ {min_to_dur_str(oft)}")
        else:
            st.markdown("**原计划调机：** 无")

    # 合并提示
    if res.get("merged_info"):
        mi = res["merged_info"]
        skipped = mi["skipped"]
        skipped_str = "、".join(
            f"{s['dep_city']}{min_to_time_str(s['dep_min'])}→"
            f"{s['arr_city']}{min_to_time_str(s['arr_min'])}"
            for s in skipped
        )
        st.info(
            f"🔗 **调机已合并：** 跳过了原调机段 `{skipped_str}`，"
            f"改为从 {mi['last_pax']['arr_city']} 直飞新需求出发地。"
        )

    if res["ferry_info"] and res["ferry_info"]["ferry_min"] > 0:
        fi = res["ferry_info"]
        tag_label = "合并调机段" if res.get("merged_info") else "新增调机段"
        st.markdown(
            f"**{tag_label}：** `{min_to_time_str(fi['dep_min'])} - "
            f"{min_to_time_str(fi['arr_min'])}`  "
            f"{fi['from_city']} → {dep_city}  "
            f"（估算飞行 {min_to_dur_str(fi['ferry_min'])}）"
        )

    if res.get("return_ferry_info"):
        rf = res["return_ferry_info"]
        st.markdown(
            f"**回程调机段：** `{min_to_time_str(rf['dep_min'])} - "
            f"{min_to_time_str(rf['arr_min'])}`  "
            f"{rf['from_city']} → {rf['to_city']}  "
            f"（接回原计划，估算飞行 {min_to_dur_str(rf['ferry_min'])}）"
        )

    if res.get("optimizations"):
        for opt in res["optimizations"]:
            st.info(
                f"💡 已自动优化 {opt['date'].strftime('%m月%d日')} 的调机段："
                f"原 `{opt['old_dep_city']} → {opt['arr_city']}`"
                f"（{opt['d_old']} km），改为 "
                f"`{opt['new_dep_city']} → {opt['arr_city']}`"
                f"（{opt['d_new']} km），省 {opt['saving_km']} km"
            )

    c1, c2 = st.columns(2)
    with c1:
        st.markdown(
            f"**当日现有值勤：** {min_to_time_str(res['duty_start'])} → "
            f"{min_to_time_str(res['duty_start'] + res['duty_total'])} "
            f"（{min_to_dur_str(res['duty_total'])}）"
        )
        st.markdown(f"**当日现有飞行：** {min_to_dur_str(res['old_flight'])}")
    with c2:
        st.markdown(
            f"**推荐主段：** `{min_to_time_str(res['main_dep'])} - "
            f"{min_to_time_str(res['main_arr'])}`  {dep_city} → {arr_city}"
        )
        st.markdown(
            f"**新增后值勤：** {min_to_time_str(res['duty_start'])} → "
            f"{min_to_time_str(res['new_duty_end'])} "
            f"（{min_to_dur_str(res['new_duty_total'])}）"
        )
        st.markdown(f"**新增后飞行：** {min_to_dur_str(res['new_flight_total'])}")

    checks = [
        f"{'✅' if res['ok_duty'] else '❌'} 值勤 ≤ 14h"
        f"（{min_to_dur_str(res['new_duty_total'])}）",
        f"{'✅' if res['ok_flight'] else '❌'} 飞行 ≤ 10h"
        f"（{min_to_dur_str(res['new_flight_total'])}）",
        f"{'✅' if res['rest_ok'] else '❌'} 前日休息 ≥ 10h"
        + (f"（{res['rest_note']}）" if res["rest_note"] else ""),
    ]
    for c in checks:
        st.markdown(f"- {c}")

    st.markdown("---")
    st.subheader("📄 可复制方案")
    st.caption("点右上角复制按钮，直接粘贴到 Jetops / 邮件 / 微信")
    st.code("\n".join(compact_lines), language=None)


# ================================================================
# 主流程
# ================================================================
if run:
    if not flights:
        st.error("请先上传航段数据 Excel"); st.stop()
    if not request_text.strip():
        st.error("请填写潜在行程请求"); st.stop()

    request = parse_request(request_text)
    missing = []
    if not request["date"]:
        missing.append("**日期**")
    if not request["route"]:
        missing.append("**航线**")
    if missing:
        st.error("❌ 需求描述中缺少：")
        for item in missing:
            st.markdown(f"- {item}")
        st.stop()

    target_date = request["date"]
    dep_city, arr_city = request["route"]
    dep_icao = get_icao(dep_city)
    arr_icao = get_icao(arr_city)

    flight_min = request["flight_min"]
    auto_estimated = False
    if flight_min is None:
        est = estimate_flight_minutes(dep_icao, arr_icao)
        if est is None:
            st.error("❌ 无法自动估算飞行时间，请手动填写。"); st.stop()
        flight_min = est
        auto_estimated = True

    plan_min = plan_minutes(flight_min)
    transit = transit_minutes(dep_icao) if dep_icao else 120
    src = flight_min_source(dep_icao, arr_icao) if auto_estimated else "用户给定"

    st.subheader("📋 请求摘要")
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.metric("日期", str(target_date))
    c2.metric("飞机", selected_reg if selected_reg else "自动选优")
    c3.metric("航线", f"{dep_icao or dep_city} → {arr_icao or arr_city}")
    c4.metric("飞行时间", min_to_dur_str(flight_min))
    c5.metric("计划时间", min_to_dur_str(plan_min))
    st.caption(
        f"飞行时间来源：**{src}** ｜ 过站：**{min_to_dur_str(transit)}**"
        f"（{'国内' if is_domestic(dep_icao) else '国际'}） ｜ "
        f"调机合并：**{'开' if merge_mode else '关'}**"
    )
    st.markdown("---")

    if auto_select:
        st.subheader("🏆 自动选优：所有飞机对比")
        with st.spinner("正在评估所有飞机…"):
            candidates = rank_aircraft_for_request(
                flights, all_regs, target_date, dep_city, arr_city,
                flight_min, relax=relax_mode, merge_ferry=merge_mode,
            )
        if not candidates:
            st.error("❌ 无法评估任何飞机"); st.stop()

        rows = []
        for c in candidates:
            if "error" in c:
                rows.append({"飞机": c["reg"], "新增/合并调机": "-", "原调机": "-",
                             "值勤": "-", "飞行": "-", "状态": f"❌ {c['error'][:36]}"})
            else:
                r = c["result"]
                new_ferry_str = (f"{c['ferry_count']} 段 / {min_to_dur_str(c['ferry_min'])}"
                                 if c["ferry_count"] else "无")
                orig_ferry_str = (f"{c['orig_ferry_count']} 段 / {min_to_dur_str(c['orig_ferry_min'])}"
                                  if c["orig_ferry_count"] else "无")
                status = "✅ 合规" if c["ok_all"] else "⚠️ 超时"
                rows.append({"飞机": c["reg"], "新增/合并调机": new_ferry_str,
                             "原调机": orig_ferry_str,
                             "值勤": min_to_dur_str(r["new_duty_total"]),
                             "飞行": min_to_dur_str(r["new_flight_total"]),
                             "状态": status})
        st.table(rows)

        valid = [c for c in candidates if "error" not in c]
        if not valid:
            st.error("❌ 没有任何飞机能完成此需求"); st.stop()

        best = valid[0]
        res_best = best["result"]
        new_ferry_str = (f"{best['ferry_count']} 段 / {min_to_dur_str(best['ferry_min'])}"
                         if best["ferry_count"] else "无需调机")
        st.success(
            f"🏆 **推荐飞机：{best['reg']}**"
            f"（调机 {new_ferry_str}，值勤 {min_to_dur_str(res_best['new_duty_total'])}，"
            f"飞行 {min_to_dur_str(res_best['new_flight_total'])}）"
        )

        st.markdown("---")
        st.subheader("📋 各飞机完整方案（按推荐顺序）")
        medals = ["🥇", "🥈", "🥉"]
        for i, c in enumerate(candidates):
            if "error" in c:
                with st.expander(f"❌ {c['reg']} —— 无法完成", expanded=False):
                    st.error(c["error"])
                continue
            rank = medals[i] if i < len(medals) else f"#{i+1}"
            r = c["result"]
            status = "✅ 合规" if c["ok_all"] else "⚠️ 超时"
            ferry_str = (f"{c['ferry_count']} 段 / {min_to_dur_str(c['ferry_min'])}"
                         if c["ferry_count"] else "无")
            header = (f"{rank} #{i+1} · {c['reg']}   ｜   调机 {ferry_str}   ｜   "
                      f"值勤 {min_to_dur_str(r['new_duty_total'])}   ｜   "
                      f"飞行 {min_to_dur_str(r['new_flight_total'])}   ｜   {status}")
            with st.expander(header, expanded=(i == 0)):
                render_result(r, c["reg"], dep_city, arr_city, compact=False)

    else:
        if not selected_reg:
            st.error("请选择飞机，或勾选「自动选最优」"); st.stop()

        if request["has_explicit_time"]:
            res = analyze_aircraft(
                flights, selected_reg, target_date, dep_city, arr_city,
                flight_min, transit, plan_min, merge_ferry=merge_mode,
            )
            if res.get("error"):
                st.error(res["error"]); st.stop()
            render_result(res, selected_reg, dep_city, arr_city)
        else:
            st.subheader("🎯 起飞时刻推荐")
            with st.spinner("正在搜索…"):
                sug = suggest_departure_times(
                    flights, selected_reg, target_date, dep_city, arr_city,
                    flight_min, objective="earliest", step=15, window_h=12,
                    relax=relax_mode, merge_ferry=merge_mode,
                )
            if sug.get("error"):
                st.error(sug["error"]); st.stop()

            baseline = sug["baseline"]
            picks = sug["picks"]
            st.info(f"📌 最早可起飞：**{min_to_time_str(baseline['main_dep'])}**")

            if sug.get("relaxed"):
                st.warning("⚠️ 无合规方案，以下为超时最少方案")
            else:
                st.success("✅ 已找到合规方案")

            table_rows = []
            for i, r in enumerate(picks):
                over_duty = max(0, r["new_duty_total"] - DUTY_MAX_MIN)
                over_flight = max(0, r["new_flight_total"] - FLIGHT_MAX_MIN)
                table_rows.append({
                    "方案": "#1 推荐" if i == 0 else f"#{i+1}",
                    "起飞": min_to_time_str(r["main_dep"]),
                    "落地": min_to_time_str(r["main_arr"]),
                    "值勤": min_to_dur_str(r["new_duty_total"]) + (
                        f" ⚠️+{min_to_dur_str(over_duty)}" if over_duty else " ✅"),
                    "飞行": min_to_dur_str(r["new_flight_total"]) + (
                        f" ⚠️+{min_to_dur_str(over_flight)}" if over_flight else " ✅"),
                    "前日休息": "✅" if r["rest_ok"] else "❌",
                })
            st.table(table_rows)

            res = picks[0]
            st.markdown("---")
            render_result(res, selected_reg, dep_city, arr_city)
